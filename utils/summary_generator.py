import pandas as pd
from datetime import datetime
from tabulate import tabulate
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
import json

def load_chart_config(path='chart_config.json'):
    with open(path, 'r') as f:
        return json.load(f)

def generate_summary(transactions):
    """
    Generates transaction summary with config-based chart settings and accurate EMI logic.
    """
    df = pd.DataFrame(transactions)

    now = datetime.now()
    timestamp = now.strftime("%Y-%m-%d_%H-%M")
    output_excel = f"output/transactions_{timestamp}.xlsx"

    # Save raw DataFrame
    df.to_excel(output_excel, index=False)

    # Load workbook to style
    wb = load_workbook(output_excel)
    ws = wb.active

    # === Color Fills ===
    light_red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    dark_red = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
    light_green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    dark_green = PatternFill(start_color='008000', end_color='008000', fill_type='solid')
    yellow = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    grey = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')

    # Step 1: Collect EMI refund credits
    emi_refunds = set()
    for row in range(2, ws.max_row + 1):
        narration = (ws[f'E{row}'].value or "").lower()
        txn_type = (ws[f'D{row}'].value or "").strip()
        amount = ws[f'C{row}'].value
        if "emi conversion" in narration and txn_type == "Credit" and amount:
            emi_refunds.add(round(float(amount), 2))

    # Step 2: Apply row color formatting
    for row in range(2, ws.max_row + 1):
        narration = (ws[f'E{row}'].value or "").upper()
        txn_type = (ws[f'D{row}'].value or "").strip()
        category = (ws[f'F{row}'].value or "").upper()
        amount = round(float(ws[f'C{row}'].value or 0), 2)

        fill = None
        if "PREVIOUS BALANCE INR" in narration:
            fill = grey
        elif "AMORTIZATION" in narration or "PROCESSING FEE" in narration or "IGST" in narration or "INTEREST" in narration:
            fill = yellow
        elif "EMI CONVERSION" in narration and txn_type == "Credit":
            fill = dark_green
        elif txn_type == "Debit" and amount in emi_refunds and category == "SHOPPING":
            fill = dark_red
        elif txn_type == "Debit":
            fill = light_red
        elif txn_type == "Credit":
            fill = light_green

        if fill:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill

    # Step 3: Filter valid rows (exclude refunded shopping)
    valid_rows = []
    for _, row in df.iterrows():
        amount = round(row["Amount"], 2)
        if row["Type"] == "Debit" and row["Category"] == "Shopping" and amount in emi_refunds:
            continue
        if row["Type"] == "Credit" and "emi conversion" in row["Narration"].lower():
            continue
        valid_rows.append(row)

    df_filtered = pd.DataFrame(valid_rows)

    # Step 4: Create category summary
    category_totals = (
        df_filtered[df_filtered["Type"] == "Debit"]
        .groupby("Category")["Amount"]
        .sum()
    )

    # Step 5: Write summary table to Excel
    summary_start_row = ws.max_row + 3
    ws.cell(row=summary_start_row, column=1).value = "===== CATEGORY SUMMARY ====="

    current_row = summary_start_row + 1
    for category, total in category_totals.items():
        ws.cell(row=current_row, column=1).value = category
        ws.cell(row=current_row, column=2).value = total
        current_row += 1

    # Step 6: Load chart config
    chart_config = load_chart_config()
    colors = chart_config.get("colors", [])
    width = chart_config.get("width", 14)
    height = chart_config.get("height", 10)
    border = chart_config.get("border", {"color": "000000", "thickness": 20000 })

    # Step 7: Add Pie Chart
    chart = PieChart()
    chart.title = "Spending Distribution"
    labels = Reference(ws, min_col=1, min_row=summary_start_row + 1, max_row=current_row - 1)
    data = Reference(ws, min_col=2, min_row=summary_start_row + 1, max_row=current_row - 1)

    chart.add_data(data, titles_from_data=False)
    chart.set_categories(labels)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showPercent = True

    chart.width = width
    chart.height = height

    if chart.graphical_properties is not None:
        if chart.graphical_properties.line is None:
            from openpyxl.drawing.line import LineProperties
            chart.graphical_properties.line = LineProperties()

        chart.graphical_properties.line.solidFill = border.get("color", "000000")
        chart.graphical_properties.line.width = border.get("thickness", 20000)

    # Apply custom slice colors
    if colors:
        for i, color in enumerate(colors):
            if i < len(chart.series[0].data_points):
                data_point = chart.series[0].data_points[i]
                if data_point.graphicalProperties is None:
                    from openpyxl.drawing.fill import FillProperties
                    data_point.graphicalProperties = FillProperties()
                data_point.graphicalProperties.solidFill = color
    # Create new worksheet for chart
    chart_ws = wb.create_sheet(title="Spending Chart")
    chart_ws.add_chart(chart, "C5")

    # Save Excel
    wb.save(output_excel)

    # Console Summary
    print("\n===== FULL TRANSACTION SUMMARY =====")
    print(tabulate(df, headers='keys', tablefmt='psql'))

    total_credit = df[df['Type'] == 'Credit']['Amount'].sum()
    total_debit = df[df['Type'] == 'Debit']['Amount'].sum()

    print("\n===== BILL SUMMARY =====")
    print(f"Total Credit : ₹{total_credit}")
    print(f"Total Debit  : ₹{total_debit}")
    print(f"Net Balance  : ₹{total_credit - total_debit}")
    print(f"\nOutput Excel saved to {output_excel}")